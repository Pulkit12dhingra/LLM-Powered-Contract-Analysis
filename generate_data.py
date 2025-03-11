import os
import random
from datetime import datetime, timedelta
from docx import Document
from fpdf import FPDF

def generate_contract_details(fraudulent=False):
    suppliers = ["Global Trade Corp", "XYZ Exports Ltd", "ABC Commodities"]
    clients = ["Metro Retailers", "Greenland Industries", "Omega Manufacturing"]
    cities = ["New York", "Los Angeles", "Houston", "Chicago", "Miami"]
    commodities = [
        {"code": "C12345", "name": "Wheat", "unit_price": 250},
        {"code": "C67890", "name": "Corn", "unit_price": 200},
        {"code": "C54321", "name": "Soybeans", "unit_price": 300}
    ]
    
    supplier = random.choice(suppliers)
    client = random.choice(clients)
    origin = random.choice(cities)
    destination = random.choice([city for city in cities if city != origin])
    commodity = random.choice(commodities)
    volume = random.randint(500, 5000)  # Metric tons
    total_price = volume * commodity['unit_price']
    
    start_date = datetime.today()
    end_date = start_date + timedelta(days=random.randint(30, 365))
    
    if fraudulent:
        # Introduce fraudulent aspects
        total_price = total_price * 10  # Inflated price
        volume = volume * 5  # Unrealistic volume
        payment_terms = "Payment to be made in advance and is non-refundable."
        termination_clause = "Contract is non-cancellable under any circumstances."
    else:
        payment_terms = "The Client agrees to make payment within 30 days upon delivery.\nLate payments will be subject to a 5% penalty."
        termination_clause = "Either party may terminate this agreement with a 30-day written notice."
    
    contract_text = f"""
    SUPPLY CHAIN CONTRACT
    ----------------------
    
    This contract is made on {start_date.strftime('%Y-%m-%d')} between {supplier} (the "Supplier") and {client} (the "Client").
    
    1. SUPPLY DETAILS:
       - Commodity: {commodity['name']} ({commodity['code']})
       - Volume: {volume} metric tons
       - Unit Price: ${commodity['unit_price']} per metric ton
       - Total Price: ${total_price}
       
    2. LOGISTICS:
       - Origin: {origin}
       - Destination: {destination}
       
    3. TERM:
       - Contract Start Date: {start_date.strftime('%Y-%m-%d')}
       - Contract End Date: {end_date.strftime('%Y-%m-%d')}
       
    4. PAYMENT TERMS:
       - {payment_terms}
       
    5. TERMINATION:
       - {termination_clause}
       
    6. FORCE MAJEURE:
       - Neither party shall be liable for delays due to unforeseen circumstances beyond their control.
       
    Signed,
    {supplier} (Supplier)        {client} (Client)
    """
    
    return contract_text, supplier, client, fraudulent

def save_contract_docx(contract_text, supplier, client, fraudulent):
    doc = Document()
    doc.add_paragraph(contract_text)
    contract_type = "fraudulent" if fraudulent else "valid"
    file_path = f"Data/sample_contracts/{supplier}_{client}_{contract_type}_contract.docx"
    doc.save(file_path)
    print(f"Contract saved as DOCX: {file_path}")

def save_contract_pdf(contract_text, supplier, client, fraudulent):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    
    for line in contract_text.split("\n"):
        pdf.cell(200, 10, txt=line, ln=True, align='L')
    
    contract_type = "fraudulent" if fraudulent else "valid"
    file_path = f"Data/sample_contracts/{supplier}_{client}_{contract_type}_contract.pdf"
    pdf.output(file_path)
    print(f"Contract saved as PDF: {file_path}")

def main():
    os.makedirs("Data/sample_contracts", exist_ok=True)
    
    # Generate a valid contract
    contract_text, supplier, client, fraudulent = generate_contract_details(False)
    save_contract_docx(contract_text, supplier, client, fraudulent)
    save_contract_pdf(contract_text, supplier, client, fraudulent)
    
    # Generate a fraudulent contract
    contract_text, supplier, client, fraudulent = generate_contract_details(True)
    save_contract_docx(contract_text, supplier, client, fraudulent)
    save_contract_pdf(contract_text, supplier, client, fraudulent)
    
    print("Sample contract generation complete!")

if __name__ == "__main__":
    main()